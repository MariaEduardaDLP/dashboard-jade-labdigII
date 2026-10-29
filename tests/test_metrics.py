from datetime import date
from io import BytesIO

from openpyxl import load_workbook
import pandas as pd
import pytest

from jade.catalog import MESSAGES, URGENCIES
from jade.data import parse_events
from jade.metrics import (
    export_table, export_xlsx, filter_events, format_duration, hourly_counts, message_counts,
    summarize, urgency_counts,
)

IDS = [message.id for message in MESSAGES]


def test_local_day_includes_late_attendance_and_utc_events(events):
    selected = filter_events(events, date(2026, 10, 1), date(2026, 10, 1), list(URGENCIES), IDS)
    assert set(selected["id"]) == {"a", "b", "c", "d"}
    result = summarize(selected)
    assert (result.total, result.attended, result.pending) == (4, 3, 1)
    assert result.average_minutes == pytest.approx(4 / 3)


def test_message_and_urgency_filters_intersect(events):
    selected = filter_events(events, date(2026, 10, 1), date(2026, 10, 2), ["Alta"], ["agua", "ajuda"])
    assert selected["id"].tolist() == ["b"]
    assert summarize(selected).average_minutes is None


def test_empty_selection_does_not_mean_all_messages(events):
    selected = filter_events(events, date(2026, 10, 1), date(2026, 10, 2), list(URGENCIES), [])
    assert summarize(selected).total == 0
    assert summarize(selected).average_minutes is None


def test_empty_api_response_supports_all_aggregations():
    empty = parse_events([])
    assert summarize(empty).pending == 0
    assert hourly_counts(empty, IDS)["count"].sum() == 0
    assert message_counts(empty, IDS)["count"].sum() == 0
    assert urgency_counts(empty)["count"].sum() == 0
    assert export_table(empty).empty


def test_hourly_series_keep_zero_hours_and_unused_messages(events):
    counts = hourly_counts(events, IDS)
    assert len(counts) == 24 * 10
    assert counts["count"].sum() == len(events)
    assert counts.query("message_id == 'comida'")["count"].eq(0).all()
    assert counts.query("message_id == 'agua' and hour == 0")["count"].item() == 1
    assert message_counts(events, IDS)["count"].sum() == len(events)
    assert urgency_counts(events)["count"].sum() == len(events)


@pytest.mark.parametrize(("minutes", "expected"), [
    (None, "—"), (0, "0 s"), (0.5, "30 s"), (2.25, "2 min 15 s"),
    (59.999, "1 h 00 min"), (90, "1 h 30 min"),
])
def test_duration_display(minutes, expected):
    assert format_duration(minutes) == expected


def test_csv_preserves_timezone_and_pending_status(events):
    table = export_table(events).set_index("Solicitação")
    assert table.loc["a", "Solicitada em"].endswith("-03:00")
    assert table.loc["b", "Atendida em"] == ""
    assert table.loc["b", "Status"] == "Pendente"
    assert table.loc["d", "Tempo de atendimento (min)"] == 0


def test_xlsx_preserves_history_timezone_pending_status_and_numeric_values(events):
    table = export_table(events)
    content = export_xlsx(table)
    restored = pd.read_excel(BytesIO(content), keep_default_na=False)
    assert restored.columns.tolist() == table.columns.tolist()
    assert restored["Solicitação"].tolist() == table["Solicitação"].tolist()
    records = restored.set_index("Solicitação")
    assert records.loc["a", "Solicitada em"].endswith("-03:00")
    assert records.loc["b", "Atendida em"] == ""
    assert records.loc["b", "Status"] == "Pendente"
    assert records.loc["d", "Tempo de atendimento (min)"] == 0
    workbook = load_workbook(BytesIO(content))
    assert workbook.active.freeze_panes == "A2"
    assert workbook.active.auto_filter.ref == "A1:G6"


def test_xlsx_handles_empty_history_and_keeps_api_identifiers_as_text():
    table = export_table(parse_events([]))
    sheet = load_workbook(BytesIO(export_xlsx(table))).active
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == table.columns.tolist()
    table.loc[0, "Solicitação"] = "=1+1"
    sheet = load_workbook(BytesIO(export_xlsx(table))).active
    assert sheet["A2"].value == "=1+1"
    assert sheet["A2"].data_type == "s"
